import json
from io import BytesIO

from django.db import transaction
from django.db.models import Count
from django.http import HttpResponse, HttpResponseBadRequest, JsonResponse
from django.shortcuts import get_object_or_404
from django.template.loader import render_to_string
from django.urls import reverse
from django.views.decorators.csrf import csrf_protect
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ...access import user_has_competicio_capability
from .base import InscripcionsListView
from ...models import Competicio, Inscripcio, InscripcioEquipAssignacio, InscripcioMedia
from ...models.competicio import CompeticioAparell, InscripcioAparellExclusio
from ...services.shared.birth_year_ranges import (
    BIRTH_YEAR_RANGE_PARTITION_CODE,
    clear_inscripcions_derived_group_config_cache,
    get_inscripcions_derived_group_config,
    normalize_birth_year_range_partition_config_for_inscripcions,
    validate_birth_year_range_partition_config,
)
from ...services.shared.competition_groups import get_group_for_display_num, get_group_maps, sync_competicio_group_names_view
from ...services.teams.equip_contexts import (
    NATIVE_EQUIP_CONTEXT_CODE,
    BASE_EQUIP_CONTEXT_NAME,
    get_contextual_assignment_map,
    get_equip_context,
    get_equip_context_summary,
    normalize_equip_context_code,
)
from ...services.scoring.team_scoring import build_team_subjects_for_comp_aparell, is_team_context_app
from ...services.teams.team_series import get_series_summary_payload
from ...services.inscripcions.history import (
    capture_inscripcions_history_snapshot,
    record_inscripcions_history_entry,
    with_inscripcions_history_payload,
)
from ...services.inscripcions.media_matching import normalize_media_matching_config
from ...services.inscripcions.admission import (
    active_baixes_qs,
    baixa_summary_by_inscripcio,
    clear_inscripcio_baixa,
    set_inscripcio_baixa,
)
from ...services.avatar.competition.overview import AVATAR_MESSAGES as COMPETITION_AVATAR_MESSAGES
from ...services.avatar.inscripcions.messages import AVATAR_MESSAGES as INSCRIPCIONS_AVATAR_MESSAGES
from ...services.inscripcions.queries import (
    _label_with_source,
    _normalize_schema_extra_code,
    _reserved_inscripcio_codes,
    build_inscripcions_sort_context_key,
    get_allowed_group_fields,
    get_available_column_filter_fields,
    get_available_sort_fields,
    get_competicio_custom_sort_codes,
    get_request_inscripcio_filters,
    reconcile_inscripcions_sort_context_state,
)
from ...services.inscripcions.timing import inscripcions_timing_section


BUILTIN_TABLE_FIELDS = [
    {"code": "nom_i_cognoms", "label": "Nom i cognoms", "kind": "builtin"},
    {"code": "document", "label": "DNI/Document", "kind": "builtin"},
    {"code": "sexe", "label": "Sexe", "kind": "builtin"},
    {"code": "data_naixement", "label": "Data naixement", "kind": "builtin"},
    {"code": "entitat", "label": "Entitat", "kind": "builtin"},
    {"code": "categoria", "label": "Categoria", "kind": "builtin"},
    {"code": "subcategoria", "label": "Subcategoria", "kind": "builtin"},
    {"code": "grup", "label": "Grup", "kind": "builtin"},
    {"code": "equip", "label": "Equip", "kind": "builtin"},
    {"code": "__aparells__", "label": "Aparells", "kind": "ui"},
    {"code": "__media__", "label": "Media", "kind": "ui"},
]

HIDDEN_TABLE_COLUMN_CODES = {"ordre_sortida"}
SYSTEM_NATIVE_TABLE_CODES = {"grup", "equip", "__aparells__", "__media__", "__actions__"}

# Frontend contract for Phase 2 incremental refresh:
# GET `inscripcions_list` accepts `__fragments=header,toolbar,history,table,panel`
# and optional `__panel_key=<panel-key>`, returning HTML snippets plus refreshed boot.
FRAGMENT_TEMPLATE_BY_NAME = {
    "header": "competicio/inscripcions/_header.html",
    "history": "competicio/inscripcions/_history_summary.html",
    "table": "competicio/inscripcions/_table.html",
    "toolbar": "competicio/inscripcions/_toolbar.html",
}

PANEL_TEMPLATE_BY_KEY = {
    "agrupacio": "competicio/inscripcions/_grouping_panel.html",
    "columnes": "competicio/inscripcions/_columns_panel.html",
    "grups": "competicio/inscripcions/_groups_panel.html",
    "equips": "competicio/inscripcions/_teams_panel.html",
    "series-equips": "competicio/inscripcions/_series_panel.html",
    "media": "competicio/inscripcions/_media_panel.html",
    "altres": "competicio/inscripcions/_altres_panel.html",
}


def _serialize_listing_media_item(item):
    return {
        "id": item.id,
        "inscripcio_id": item.inscripcio_id,
        "tipus": item.tipus,
        "mime_type": item.mime_type or "",
        "original_filename": item.original_filename or "",
        "file_size_bytes": int(item.file_size_bytes or 0),
        "is_primary": bool(item.is_primary),
        "source": item.source or "",
        "match_score": float(item.match_score) if item.match_score is not None else None,
        "url": reverse(
            "inscripcions_media_file",
            kwargs={"pk": item.competicio_id, "media_id": item.id},
        ),
    }


def _build_url_template(url_name, *, pk, placeholder_kwargs):
    kwargs = {"pk": pk}
    replacements = []
    for key, config in (placeholder_kwargs or {}).items():
        dummy = config["dummy"]
        placeholder = config["placeholder"]
        kwargs[key] = dummy
        replacements.append((str(dummy), placeholder))

    url = reverse(url_name, kwargs=kwargs)
    for dummy, placeholder in replacements:
        url = url.replace(dummy, placeholder, 1)
    return url


def get_available_table_columns(competicio):
    out = []
    seen = set()
    reserved = _reserved_inscripcio_codes()

    schema = competicio.inscripcions_schema or {}
    columns = schema.get("columns") or []
    excel_codes = set()
    if isinstance(columns, list):
        for col in columns:
            if not isinstance(col, dict):
                continue
            code = col.get("code")
            if not code:
                continue
            kind = col.get("kind") or "extra"
            if kind == "extra":
                code = _normalize_schema_extra_code(code, reserved)
            if code in HIDDEN_TABLE_COLUMN_CODES:
                continue
            excel_codes.add(code)

    for field in BUILTIN_TABLE_FIELDS:
        if field["code"] in seen:
            continue
        source = "native" if field["code"] in SYSTEM_NATIVE_TABLE_CODES else ("excel" if field["code"] in excel_codes else "native")
        out.append({**field, "source": source, "ui_label": _label_with_source(field["label"], source)})
        seen.add(field["code"])

    if isinstance(columns, list):
        for col in columns:
            if not isinstance(col, dict):
                continue
            code = col.get("code")
            if not code:
                continue
            kind = col.get("kind") or "extra"
            if kind == "extra":
                code = _normalize_schema_extra_code(code, reserved)
            if code in HIDDEN_TABLE_COLUMN_CODES:
                continue
            if code in seen:
                continue
            label = col.get("label") or code
            out.append(
                {
                    "code": code,
                    "label": label,
                    "kind": kind,
                    "source": "excel",
                    "ui_label": _label_with_source(label, "excel"),
                }
            )
            seen.add(code)

    out.append(
        {
            "code": "__actions__",
            "label": "Accions",
            "kind": "ui",
            "source": "native",
            "ui_label": _label_with_source("Accions", "native"),
        }
    )
    return out


def get_selected_table_columns(competicio, available_cols):
    view_cfg = competicio.inscripcions_view or {}
    selected_codes = view_cfg.get("table_columns")
    if not isinstance(selected_codes, list) or not selected_codes:
        selected_codes = [
            "nom_i_cognoms",
            "document",
            "sexe",
            "data_naixement",
            "entitat",
            "categoria",
            "subcategoria",
            "grup",
            "equip",
            "__aparells__",
            "__media__",
            "__actions__",
        ]

    by_code = {col["code"]: col for col in available_cols}
    selected = [by_code[code] for code in selected_codes if code in by_code]
    if not selected and "nom_i_cognoms" in by_code:
        selected = [by_code["nom_i_cognoms"]]
    return selected


class InscripcionsListNewView(InscripcionsListView):
    template_name = "competicio/inscripcions/inscripcions_page.html"
    enable_lazy_group_tabs = True

    def _get_requested_fragments(self):
        raw_value = str(self.request.GET.get("__fragments") or "").strip()
        if not raw_value:
            return []
        out = []
        for token in raw_value.split(","):
            name = str(token or "").strip().lower()
            if name in FRAGMENT_TEMPLATE_BY_NAME and name not in out:
                out.append(name)
            elif name == "panel" and name not in out:
                out.append(name)
        return out

    def _render_fragment_payload(self, context):
        requested_fragments = self._get_requested_fragments()
        if not requested_fragments:
            return None

        payload = {
            "ok": True,
            "fragments": {},
            "boot": context.get("inscripcions_page_boot") or {},
        }
        for name in requested_fragments:
            if name == "panel":
                panel_key = str(self.request.GET.get("__panel_key") or "").strip()
                template_name = PANEL_TEMPLATE_BY_KEY.get(panel_key)
                if not template_name:
                    continue
                payload["fragments"]["panel"] = {
                    "panel_key": panel_key,
                    "html": render_to_string(template_name, context, request=self.request),
                }
                continue

            template_name = FRAGMENT_TEMPLATE_BY_NAME.get(name)
            if not template_name:
                continue
            payload["fragments"][name] = render_to_string(template_name, context, request=self.request)
        return payload

    def render_to_response(self, context, **response_kwargs):
        fragment_payload = self._render_fragment_payload(context)
        if fragment_payload is not None:
            return JsonResponse(fragment_payload)
        return super().render_to_response(context, **response_kwargs)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        requested_fragments = self._get_requested_fragments()
        requested_panel_key = str(self.request.GET.get("__panel_key") or "").strip()
        ctx["inscripcions_lazy_panels"] = not ("panel" in requested_fragments and requested_panel_key)
        ctx["avatar_messages"] = {**COMPETITION_AVATAR_MESSAGES, **INSCRIPCIONS_AVATAR_MESSAGES}
        ctx["avatar_initial_topic"] = "competition_inscriptions"
        materialized_records = list(ctx.get("_inscripcions_materialized_records") or [])
        can_reuse_materialized = bool(materialized_records) and not bool(ctx.get("is_paginated"))
        with inscripcions_timing_section(self.request, "listing.permissions_counts"):
            ctx["can_edit_inscripcions"] = user_has_competicio_capability(
                self.request.user,
                self.competicio,
                "inscripcions.edit",
            )
            filtered_qs = None
            if can_reuse_materialized:
                ctx["inscrits_filtered_count"] = len(materialized_records)
                ctx["existing_groups_count"] = len(
                    {
                        int(group_num)
                        for group_num in (
                            getattr(row, "grup", None)
                            for row in materialized_records
                        )
                        if group_num is not None
                    }
                )
            else:
                filtered_qs = self.get_queryset_base_filtrada()
                ctx["inscrits_filtered_count"] = filtered_qs.count()
                ctx["existing_groups_count"] = filtered_qs.exclude(grup__isnull=True).values("grup").distinct().count()
            ctx["inscrits_total_count"] = Inscripcio.objects.filter(competicio=self.competicio).count()

        with inscripcions_timing_section(self.request, "listing.table_context"):
            available_table_columns = get_available_table_columns(self.competicio)
            selected_table_columns = get_selected_table_columns(self.competicio, available_table_columns)
            active_filters = get_request_inscripcio_filters(self.request, competicio=self.competicio)
            column_filter_fields = get_available_column_filter_fields(self.competicio)
            filterable_codes = {field["code"] for field in column_filter_fields}
            active_column_filters = dict(active_filters.get("column_filters") or {})
            selected_table_column_codes = {
                col["code"]
                for col in selected_table_columns
                if isinstance(col, dict) and col.get("code")
            }

            ctx["available_table_columns"] = available_table_columns
            ctx["selected_table_columns"] = selected_table_columns
            ctx["table_colspan"] = len(selected_table_columns) if selected_table_columns else 1
            sort_options = ctx.get("sort_field_options") or []
            ctx["sortable_table_column_codes"] = [item.get("code") for item in sort_options if isinstance(item, dict) and item.get("code")]
            sortable_codes = set(ctx["sortable_table_column_codes"])
            ctx["filterable_table_column_codes"] = sorted(filterable_codes)
            ctx["column_menu_codes"] = sorted(sortable_codes | filterable_codes)
            custom_sort_codes = set(get_competicio_custom_sort_codes(self.competicio, allowed_sort_codes=sortable_codes))

        with inscripcions_timing_section(self.request, "listing.sort_context"):
            active_group_by = list(ctx.get("selected_group_fields") or [])
            sort_context_key = build_inscripcions_sort_context_key(self.competicio.id, filters=active_filters, group_by=active_group_by)
            if can_reuse_materialized:
                current_ids = [int(row.id) for row in materialized_records if getattr(row, "id", None) is not None]
            else:
                if filtered_qs is None:
                    filtered_qs = self.get_queryset_base_filtrada()
                current_ids = list(filtered_qs.order_by("ordre_sortida", "id").values_list("id", flat=True))
            sort_state = reconcile_inscripcions_sort_context_state(self.request, sort_context_key, current_ids)
            sort_stack = [entry for entry in (sort_state.get("stack") or []) if isinstance(entry, dict) and entry.get("sort_key") in sortable_codes]
            competition_order_tail_active = bool(sort_stack) and bool(sort_state.get("competition_order_tail"))

        with inscripcions_timing_section(self.request, "listing.sort_entries"):
            sort_label_by_code = {}
            for item in sort_options:
                if not isinstance(item, dict):
                    continue
                code = item.get("code")
                if not code:
                    continue
                sort_label_by_code[code] = item.get("ui_label") or item.get("label") or code

            dir_to_symbol = {"asc": "\u2191", "desc": "\u2193", "arrow_asc": "\u2195\u2191", "arrow_desc": "\u2195\u2193", "custom": "P"}
            dir_to_label = {
                "asc": "Ascendent",
                "desc": "Descendent",
                "arrow_asc": "Fletxa ascendent",
                "arrow_desc": "Fletxa descendent",
                "custom": "Personalitzat",
            }

            sort_entries = []
            for priority, entry in enumerate(sort_stack, start=1):
                sort_code = str(entry.get("sort_key") or "")
                if not sort_code:
                    continue
                sort_dir = str(entry.get("sort_dir") or "asc")
                scope = str(entry.get("scope") or "all")
                scope_short = "TF"
                scope_label = "Totes les inscripcions filtrades"
                if scope == "tab":
                    scope_short = "P"
                    scope_label = "Dins de cada pestanya activa"
                elif scope == "all_groups":
                    scope_short = "TG"
                    scope_label = "Dins de cada grup numeric complet"
                elif scope == "group":
                    group_num = entry.get("group_num")
                    scope_short = f"G{group_num}" if group_num else "G"
                    scope_label = f"Nomes grup numeric {group_num} complet" if group_num else "Nomes un grup numeric concret"

                symbol = dir_to_symbol.get(sort_dir, "\u2191")
                sort_entries.append(
                    {
                        "priority": priority,
                        "remove_priority": priority,
                        "code": sort_code,
                        "label": sort_label_by_code.get(sort_code, sort_code),
                        "symbol": symbol,
                        "sort_dir": sort_dir,
                        "sort_dir_label": dir_to_label.get(sort_dir, sort_dir),
                        "scope": scope,
                        "scope_short": scope_short,
                        "scope_label": scope_label,
                        "custom_active": sort_dir == "custom",
                        "title": f"#{priority} - {dir_to_label.get(sort_dir, sort_dir)} - {scope_label}",
                    }
                )

            indicator_by_code = {}
            for item in sort_entries:
                sort_code = item["code"]
                if sort_code in indicator_by_code:
                    continue
                indicator_by_code[sort_code] = {
                    "priority": item["priority"],
                    "remove_priority": item["remove_priority"],
                    "symbol": item["symbol"],
                    "scope_short": item["scope_short"],
                    "title": item["title"],
                }

            ctx["column_sort_context_key"] = sort_context_key
            ctx["column_sort_stack"] = sort_stack
            ctx["column_sort_has_stack"] = bool(sort_stack)
            ctx["column_sort_stack_count"] = len(sort_stack)
            ctx["column_sort_effective_count"] = len(sort_stack) + (1 if competition_order_tail_active else 0)
            ctx["column_sort_competition_tail_active"] = competition_order_tail_active
            ctx["column_sort_competition_tail_priority"] = len(sort_stack) + 1
            ctx["column_sort_entries"] = sort_entries
            ctx["column_sort_indicator_by_code"] = indicator_by_code
            ctx["custom_sort_enabled_by_code"] = {code: True for code in custom_sort_codes}
            ctx["column_filter_indicator_by_code"] = {
                code: {"count": len(tokens), "title": f"{len(tokens)} valor(s) filtrat(s)"}
                for code, tokens in active_column_filters.items()
                if tokens
            }
            ctx["column_filter_tokens_by_code"] = active_column_filters
            ctx["active_column_filter_items"] = [{"param": f"cf_{code}", "token": token} for code, tokens in active_column_filters.items() for token in tokens]

        with inscripcions_timing_section(self.request, "listing.workspace_context"):
            group_field_options = get_allowed_group_fields(self.competicio)
            derived_group_cfg = get_inscripcions_derived_group_config(self.competicio.inscripcions_view or {})
            ctx["birth_year_range_group_config"] = derived_group_cfg.get(BIRTH_YEAR_RANGE_PARTITION_CODE) or {"ranges": []}

            ctx["group_names"] = get_group_maps(self.competicio).get("name_map") or {}
            visible_group_nums = set()
            records_grouped = ctx.get("records_grouped")
            if records_grouped:
                for _label, rows, _group_key in records_grouped:
                    for row in rows:
                        group_num = getattr(row, "grup", None)
                        if group_num is not None:
                            visible_group_nums.add(int(group_num))
            else:
                for row in (ctx.get("records") or []):
                    group_num = getattr(row, "grup", None)
                    if group_num is not None:
                        visible_group_nums.add(int(group_num))

            group_totals_qs = Inscripcio.objects.filter(competicio=self.competicio).exclude(grup__isnull=True)
            if visible_group_nums:
                group_totals_qs = group_totals_qs.filter(grup__in=visible_group_nums)
            ctx["group_member_totals"] = {
                str(row["grup"]): int(row["total"] or 0)
                for row in group_totals_qs.values("grup").annotate(total=Count("id"))
                if row.get("grup") is not None
            }

            team_context_code = normalize_equip_context_code(self.request.GET.get("team_context"))
            selected_team_context = get_equip_context(self.competicio, team_context_code)
            if selected_team_context is None:
                team_context_code = NATIVE_EQUIP_CONTEXT_CODE
                selected_team_context = get_equip_context(self.competicio, team_context_code)
            team_fields = group_field_options
            ctx["team_partition_fields"] = team_fields
            ctx["team_context_selected_code"] = team_context_code
            ctx["team_context_selected_label"] = str(
                getattr(selected_team_context, "nom", "") or (
                    "Base" if team_context_code == NATIVE_EQUIP_CONTEXT_CODE else team_context_code
                )
            ).strip() or team_context_code
            aparells_cfg = list(
                CompeticioAparell.objects.filter(competicio=self.competicio, actiu=True).select_related("aparell").order_by("ordre", "id")
            )
            individual_aparells_cfg = [app for app in aparells_cfg if not is_team_context_app(app)]
            active_app_ids = [app.id for app in individual_aparells_cfg]
            ctx["inscripcio_aparells_cfg"] = individual_aparells_cfg
            ctx["inscripcio_aparells_active_ids"] = active_app_ids
            ctx["sidebar_team_nav_count"] = int(
                get_equip_context_summary(self.competicio, team_context_code).get("teams_total") or 0
            )
            series_team_aparells = [app for app in aparells_cfg if is_team_context_app(app)]
            ctx["series_team_aparells"] = series_team_aparells
            ctx["series_default_comp_aparell"] = series_team_aparells[0] if series_team_aparells else None
            default_series_comp_aparell = series_team_aparells[0] if series_team_aparells else None
            if default_series_comp_aparell is not None:
                series_subjects, _series_issues = build_team_subjects_for_comp_aparell(
                    self.competicio,
                    default_series_comp_aparell,
                )
                ctx["sidebar_series_nav_count"] = int(
                    get_series_summary_payload(
                        self.competicio,
                        default_series_comp_aparell,
                        series_subjects,
                    ).get("series_total") or 0
                )
            else:
                ctx["sidebar_series_nav_count"] = 0

        with inscripcions_timing_section(self.request, "listing.media_maps"):
            visible_ins_ids = set()
            visible_rows = []
            table_runtime = {}
            records_grouped = ctx.get("records_grouped")
            if records_grouped:
                for _label, rows, _group_key in records_grouped:
                    for row in rows:
                        visible_rows.append(row)
                        base_equip_id = getattr(row, "_base_equip_id_cache", None)
                        base_equip_name = getattr(row, "_base_equip_name_cache", "")
                        setattr(row, "base_equip_id", base_equip_id)
                        setattr(row, "base_equip_name", base_equip_name)
                        if getattr(row, "id", None):
                            visible_ins_ids.add(row.id)
            else:
                for row in (ctx.get("records") or []):
                    visible_rows.append(row)
                    base_equip_id = getattr(row, "_base_equip_id_cache", None)
                    base_equip_name = getattr(row, "_base_equip_name_cache", "")
                    setattr(row, "base_equip_id", base_equip_id)
                    setattr(row, "base_equip_name", base_equip_name)
                    if getattr(row, "id", None):
                        visible_ins_ids.add(row.id)

            if "equip" in selected_table_column_codes:
                legacy_team_by_ins_id = {}
                if visible_ins_ids:
                    legacy_rows = (
                        Inscripcio.objects
                        .filter(id__in=visible_ins_ids)
                        .select_related("equip")
                        .values("id", "equip_id", "equip__nom")
                    )
                    for item in legacy_rows:
                        equip_id = item.get("equip_id")
                        if not equip_id:
                            continue
                        legacy_team_by_ins_id[int(item["id"])] = {
                            "id": equip_id,
                            "name": str(item.get("equip__nom") or "").strip(),
                        }
                selected_assignment_map = (
                    get_contextual_assignment_map(self.competicio, visible_ins_ids, team_context_code)
                    if visible_ins_ids
                    else {}
                )
                assignments_by_ins_id = {}
                if visible_ins_ids:
                    assignment_rows = (
                        InscripcioEquipAssignacio.objects
                        .filter(competicio=self.competicio, inscripcio_id__in=visible_ins_ids)
                        .select_related("context", "equip")
                        .order_by("context__nom", "equip__nom", "id")
                    )
                    for assignment in assignment_rows:
                        context = getattr(assignment, "context", None)
                        equip = getattr(assignment, "equip", None)
                        context_code = str(getattr(context, "code", "") or "").strip()
                        context_label = (
                            BASE_EQUIP_CONTEXT_NAME
                            if context_code == NATIVE_EQUIP_CONTEXT_CODE
                            else str(getattr(context, "nom", "") or "").strip()
                        ) or context_code
                        equip_name = str(getattr(equip, "nom", "") or "").strip()
                        if not context_label or not equip_name:
                            continue
                        assignments_by_ins_id.setdefault(int(assignment.inscripcio_id), []).append(
                            {
                                "context_code": context_code,
                                "context_label": context_label,
                                "equip_name": equip_name,
                            }
                        )
                for ins_id, legacy_team in legacy_team_by_ins_id.items():
                    existing_native = any(
                        item.get("context_code") == NATIVE_EQUIP_CONTEXT_CODE
                        for item in assignments_by_ins_id.get(ins_id, [])
                    )
                    if existing_native:
                        continue
                    assignments_by_ins_id.setdefault(ins_id, []).append(
                        {
                            "context_code": NATIVE_EQUIP_CONTEXT_CODE,
                            "context_label": BASE_EQUIP_CONTEXT_NAME,
                            "equip_name": legacy_team.get("name") or f"Equip {legacy_team.get('id')}",
                        }
                    )
                selected_context_label = str(ctx.get("team_context_selected_label") or team_context_code).strip()
                for row in visible_rows:
                    row_id = getattr(row, "id", None)
                    selected_assignment = selected_assignment_map.get(row_id)
                    selected_equip = getattr(selected_assignment, "equip", None) if selected_assignment is not None else None
                    if selected_equip is None and team_context_code == NATIVE_EQUIP_CONTEXT_CODE:
                        legacy_team = legacy_team_by_ins_id.get(int(row_id or 0), {})
                        selected_equip_id = getattr(row, "base_equip_id", None) or legacy_team.get("id")
                        selected_equip_name = getattr(row, "base_equip_name", "") or legacy_team.get("name") or ""
                    else:
                        selected_equip_id = getattr(selected_equip, "id", None)
                        selected_equip_name = str(getattr(selected_equip, "nom", "") or "").strip()
                    all_assignments = assignments_by_ins_id.get(int(row_id or 0), [])
                    other_assignments = [
                        item
                        for item in all_assignments
                        if item.get("context_code") != team_context_code
                    ]
                    tooltip_lines = [f"Context actual: {selected_context_label}"]
                    if selected_equip_name:
                        tooltip_lines.append(f"Equip actual: {selected_equip_name}")
                    else:
                        tooltip_lines.append("Sense equip en aquest context")
                    if other_assignments:
                        tooltip_lines.append("Altres contextos:")
                        tooltip_lines.extend(
                            f"{item['context_label']}: {item['equip_name']}"
                            for item in other_assignments[:6]
                        )
                        if len(other_assignments) > 6:
                            tooltip_lines.append(f"+{len(other_assignments) - 6} contextos mes")
                    setattr(row, "context_equip_id", selected_equip_id)
                    setattr(row, "context_equip_name", selected_equip_name)
                    setattr(row, "team_context_label", selected_context_label)
                    setattr(row, "other_team_contexts_count", len(other_assignments))
                    setattr(row, "team_context_tooltip", "\n".join(tooltip_lines))

            if "__aparells__" in selected_table_column_codes:
                excluded_map = {}
                if visible_ins_ids and active_app_ids:
                    excl_pairs = InscripcioAparellExclusio.objects.filter(
                        inscripcio_id__in=visible_ins_ids,
                        comp_aparell_id__in=active_app_ids,
                    ).values_list("inscripcio_id", "comp_aparell_id")
                    for ins_id, app_id in excl_pairs:
                        excluded_map.setdefault(str(ins_id), []).append(app_id)
                    for ins_id in excluded_map:
                        excluded_map[ins_id].sort()
                for ins_id in visible_ins_ids:
                    excluded_map.setdefault(str(ins_id), [])
                table_runtime["inscripcio_aparells_excluded_map"] = excluded_map

            if visible_ins_ids:
                table_runtime["inscripcio_baixes_map"] = baixa_summary_by_inscripcio(
                    self.competicio,
                    active_app_ids,
                    inscripcio_ids=visible_ins_ids,
                )

            if "__media__" in selected_table_column_codes:
                media_map = {}
                if visible_ins_ids:
                    media_qs = (
                        InscripcioMedia.objects.filter(competicio=self.competicio, inscripcio_id__in=visible_ins_ids)
                        .order_by("inscripcio_id", "-is_primary", "-created_at", "id")
                    )
                    for media in media_qs:
                        media_map.setdefault(str(media.inscripcio_id), []).append(_serialize_listing_media_item(media))
                table_runtime["inscripcio_media_map"] = media_map
            if table_runtime:
                ctx["table_runtime"] = table_runtime
            ctx["inscripcions_page_boot"] = {
                "ids": {
                    "competicioId": self.competicio.id,
                },
                "flags": {
                    "canEdit": bool(ctx["can_edit_inscripcions"]),
                    "sortContextKey": sort_context_key,
                    "teamContextSelectedCode": team_context_code,
                },
                "urls": {
                    "historyUndo": reverse("inscripcions_history_undo", kwargs={"pk": self.competicio.id}),
                    "historyRedo": reverse("inscripcions_history_redo", kwargs={"pk": self.competicio.id}),
                    "sortApply": reverse("inscripcions_sort_apply", kwargs={"pk": self.competicio.id}),
                    "sortRemove": reverse("inscripcions_sort_remove", kwargs={"pk": self.competicio.id}),
                    "sortClear": reverse("inscripcions_sort_clear", kwargs={"pk": self.competicio.id}),
                    "sortTailToggle": reverse("inscripcions_sort_competition_tail_toggle", kwargs={"pk": self.competicio.id}),
                    "filterValues": reverse("inscripcions_filter_values", kwargs={"pk": self.competicio.id}),
                    "sortCustomValues": reverse("inscripcions_sort_custom_values", kwargs={"pk": self.competicio.id}),
                    "sortCustomSave": reverse("inscripcions_sort_custom_save", kwargs={"pk": self.competicio.id}),
                    "saveBirthYearRangeConfig": reverse("inscripcions_save_birth_year_range_config", kwargs={"pk": self.competicio.id}),
                    "mediaUpload": reverse("inscripcions_media_upload", kwargs={"pk": self.competicio.id}),
                    "mediaSetPrimary": reverse("inscripcions_media_set_primary", kwargs={"pk": self.competicio.id}),
                    "mediaDelete": reverse("inscripcions_media_delete", kwargs={"pk": self.competicio.id}),
                    "mediaSaveMatchingConfig": reverse("inscripcions_media_match_config_save", kwargs={"pk": self.competicio.id}),
                    "mediaMatchPreview": reverse("inscripcions_media_match_preview", kwargs={"pk": self.competicio.id}),
                    "mediaMatchApply": reverse("inscripcions_media_match_apply", kwargs={"pk": self.competicio.id}),
                    "mediaWorkspace": reverse("inscripcions_media_workspace", kwargs={"pk": self.competicio.id}),
                    "mediaReassign": reverse("inscripcions_media_reassign", kwargs={"pk": self.competicio.id}),
                    "reorder": reverse("inscripcions_reorder", kwargs={"pk": self.competicio.id}),
                    "saveTableColumns": reverse("inscripcions_save_table_columns", kwargs={"pk": self.competicio.id}),
                    "setGroupName": reverse("inscripcions_set_group_name", kwargs={"pk": self.competicio.id}),
                    "setAparells": reverse("inscripcions_set_aparells", kwargs={"pk": self.competicio.id}),
                    "setBaixa": reverse("inscripcions_set_baixa", kwargs={"pk": self.competicio.id}),
                    "clearBaixa": reverse("inscripcions_clear_baixa", kwargs={"pk": self.competicio.id}),
                    "exportBaixes": reverse("inscripcions_baixes_export", kwargs={"pk": self.competicio.id}),
                    "mergeTabs": reverse("inscripcions_merge_tabs", kwargs={"pk": self.competicio.id}),
                    "groupCompetitionOrderPreview": reverse("inscripcions_group_competition_order_preview", kwargs={"pk": self.competicio.id}),
                    "saveGroupCompetitionOrder": reverse("inscripcions_save_group_competition_order", kwargs={"pk": self.competicio.id}),
                    "bulkGroupCompetitionOrderPreview": reverse("inscripcions_bulk_group_competition_order_preview", kwargs={"pk": self.competicio.id}),
                    "bulkGroupCompetitionOrderApply": reverse("inscripcions_bulk_group_competition_order_apply", kwargs={"pk": self.competicio.id}),
                },
                "initial": {
                    "historyState": ctx["history_state"],
                    "birthYearRangeGroupConfig": ctx["birth_year_range_group_config"],
                    "mediaMatchingConfig": normalize_media_matching_config(
                        (self.competicio.inscripcions_view or {}).get("media_matching")
                    ),
                    "baixaAparells": [
                        {"id": int(app.id), "label": str(getattr(app, "display_nom", "") or app.id)}
                        for app in individual_aparells_cfg
                    ],
                },
            }
        return ctx


@require_POST
@csrf_protect
def inscripcions_save_table_columns(request, pk):
    competicio = get_object_or_404(Competicio, pk=pk)
    try:
        payload = json.loads(request.body.decode("utf-8"))
    except Exception:
        return HttpResponseBadRequest("JSON invalid")

    cols = payload.get("table_columns") or payload.get("columns") or []
    if not isinstance(cols, list):
        return HttpResponseBadRequest("table_columns ha de ser una llista")

    available = get_available_table_columns(competicio)
    allowed_codes = {item["code"] for item in available}
    cleaned = []
    for value in cols:
        if not isinstance(value, str):
            continue
        if value in allowed_codes and value not in cleaned:
            cleaned.append(value)
    if not cleaned:
        return HttpResponseBadRequest("No hi ha cap columna valida")

    before_snapshot = capture_inscripcions_history_snapshot(request, competicio)
    view_cfg = competicio.inscripcions_view or {}
    view_cfg["table_columns"] = cleaned
    competicio.inscripcions_view = view_cfg
    competicio.save(update_fields=["inscripcions_view"])
    record_inscripcions_history_entry(
        request,
        competicio,
        action_type="save_table_columns",
        action_label="Desar columnes de taula",
        before_snapshot=before_snapshot,
        after_snapshot=capture_inscripcions_history_snapshot(request, competicio),
    )
    return JsonResponse(with_inscripcions_history_payload({"ok": True, "table_columns": cleaned}, request, competicio.id))


@require_POST
@csrf_protect
def inscripcions_set_group_name(request, pk):
    competicio = get_object_or_404(Competicio, pk=pk)
    try:
        payload = json.loads(request.body.decode("utf-8"))
    except Exception:
        return HttpResponseBadRequest("JSON invalid")

    group = payload.get("group")
    name = str(payload.get("name") or "").strip()
    try:
        group_int = int(group)
    except Exception:
        return HttpResponseBadRequest("group invalid")

    before_snapshot = capture_inscripcions_history_snapshot(request, competicio)
    group_obj = get_group_for_display_num(competicio, group_int)
    if group_obj is None:
        return HttpResponseBadRequest("group invalid")
    if group_obj.nom != name:
        group_obj.nom = name
        group_obj.save(update_fields=["nom"])
    sync_competicio_group_names_view(competicio)
    record_inscripcions_history_entry(
        request,
        competicio,
        action_type="set_group_name",
        action_label="Desar nom de grup",
        before_snapshot=before_snapshot,
        after_snapshot=capture_inscripcions_history_snapshot(request, competicio),
    )
    return JsonResponse(with_inscripcions_history_payload({"ok": True, "group": group_int, "name": name}, request, competicio.id))


@require_POST
@csrf_protect
def inscripcions_save_birth_year_range_config(request, pk):
    competicio = get_object_or_404(Competicio, pk=pk)
    try:
        payload = json.loads(request.body.decode("utf-8"))
    except Exception:
        return HttpResponseBadRequest("JSON invalid")

    raw_cfg = payload.get("config") if isinstance(payload, dict) else {}
    cfg, errors = validate_birth_year_range_partition_config(raw_cfg, require_ranges=True)
    if errors:
        return HttpResponseBadRequest("\n".join(errors))
    storage_cfg = normalize_birth_year_range_partition_config_for_inscripcions(cfg)

    before_snapshot = capture_inscripcions_history_snapshot(request, competicio)
    view_cfg = dict(competicio.inscripcions_view or {})
    derived_group_cfg = view_cfg.get("derived_group_config")
    if not isinstance(derived_group_cfg, dict):
        derived_group_cfg = {}
    else:
        derived_group_cfg = dict(derived_group_cfg)
    derived_group_cfg[BIRTH_YEAR_RANGE_PARTITION_CODE] = storage_cfg
    view_cfg["derived_group_config"] = derived_group_cfg
    competicio.inscripcions_view = view_cfg
    competicio.save(update_fields=["inscripcions_view"])
    clear_inscripcions_derived_group_config_cache()
    record_inscripcions_history_entry(
        request,
        competicio,
        action_type="save_birth_year_range_group_config",
        action_label="Desar forquilles de data de naixement",
        before_snapshot=before_snapshot,
        after_snapshot=capture_inscripcions_history_snapshot(request, competicio),
    )
    return JsonResponse(with_inscripcions_history_payload({"ok": True, "config": storage_cfg}, request, competicio.id))


@require_POST
@csrf_protect
def inscripcions_set_aparells(request, pk):
    competicio = get_object_or_404(Competicio, pk=pk)
    try:
        payload = json.loads(request.body.decode("utf-8"))
    except Exception:
        return HttpResponseBadRequest("JSON invalid")

    inscripcio_id = payload.get("inscripcio_id")
    selected_ids_raw = payload.get("selected_comp_aparell_ids") or []
    try:
        inscripcio_id = int(inscripcio_id)
    except Exception:
        return HttpResponseBadRequest("inscripcio_id invalid")
    if not isinstance(selected_ids_raw, list):
        return HttpResponseBadRequest("selected_comp_aparell_ids ha de ser una llista")

    inscripcio = get_object_or_404(Inscripcio, pk=inscripcio_id, competicio=competicio)
    active_ids = [
        int(comp_aparell.id)
        for comp_aparell in (
            CompeticioAparell.objects
            .filter(competicio=competicio, actiu=True)
            .select_related("aparell")
            .order_by("ordre", "id")
        )
        if not is_team_context_app(comp_aparell)
    ]
    active_set = set(active_ids)
    selected_set = set()
    for value in selected_ids_raw:
        try:
            clean = int(value)
        except Exception:
            return HttpResponseBadRequest("selected_comp_aparell_ids conte valors invalids")
        if clean not in active_set:
            return HttpResponseBadRequest("selected_comp_aparell_ids conte aparells no valids per la competicio")
        selected_set.add(clean)

    excluded_ids = [app_id for app_id in active_ids if app_id not in selected_set]
    before_snapshot = capture_inscripcions_history_snapshot(request, competicio)
    with transaction.atomic():
        InscripcioAparellExclusio.objects.filter(inscripcio=inscripcio, comp_aparell_id__in=active_ids).delete()
        if excluded_ids:
            InscripcioAparellExclusio.objects.bulk_create(
                [InscripcioAparellExclusio(inscripcio_id=inscripcio.id, comp_aparell_id=app_id) for app_id in excluded_ids]
            )

    selected_ids = [app_id for app_id in active_ids if app_id in selected_set]
    record_inscripcions_history_entry(
        request,
        competicio,
        action_type="set_aparells",
        action_label="Desar aparells de la inscripcio",
        before_snapshot=before_snapshot,
        after_snapshot=capture_inscripcions_history_snapshot(request, competicio),
    )
    return JsonResponse(
        with_inscripcions_history_payload(
            {
                "ok": True,
                "inscripcio_id": inscripcio.id,
                "active_comp_aparell_ids": active_ids,
                "selected_comp_aparell_ids": selected_ids,
                "excluded_comp_aparell_ids": excluded_ids,
            },
            request,
            competicio.id,
        )
    )


@require_POST
@csrf_protect
def inscripcions_set_baixa(request, pk):
    competicio = get_object_or_404(Competicio, pk=pk)
    try:
        payload = json.loads(request.body.decode("utf-8"))
    except Exception:
        return HttpResponseBadRequest("JSON invalid")

    try:
        inscripcio_id = int(payload.get("inscripcio_id"))
    except Exception:
        return HttpResponseBadRequest("inscripcio_id invalid")

    inscripcio = get_object_or_404(Inscripcio, pk=inscripcio_id, competicio=competicio)
    scope = str(payload.get("scope") or "").strip().lower()
    global_scope = scope in {"global", "all", "tota", "tot"}
    app_ids_raw = payload.get("comp_aparell_ids") or payload.get("app_ids") or []
    if not isinstance(app_ids_raw, list):
        return HttpResponseBadRequest("comp_aparell_ids ha de ser una llista")
    if not global_scope and not app_ids_raw:
        return HttpResponseBadRequest("Cal indicar almenys un aparell o marcar baixa global")

    before_snapshot = capture_inscripcions_history_snapshot(request, competicio)
    try:
        set_inscripcio_baixa(
            competicio,
            inscripcio,
            app_ids=app_ids_raw,
            global_scope=global_scope,
            motiu=payload.get("motiu") or "",
            notes=payload.get("notes") or "",
            user=request.user if getattr(request.user, "is_authenticated", False) else None,
        )
    except ValueError as exc:
        return HttpResponseBadRequest(str(exc))

    record_inscripcions_history_entry(
        request,
        competicio,
        action_type="set_baixa",
        action_label="Marcar baixa",
        before_snapshot=before_snapshot,
        after_snapshot=capture_inscripcions_history_snapshot(request, competicio),
    )
    summary = baixa_summary_by_inscripcio(competicio, inscripcio_ids=[inscripcio.id]).get(str(inscripcio.id), {})
    return JsonResponse(
        with_inscripcions_history_payload(
            {
                "ok": True,
                "inscripcio_id": inscripcio.id,
                "baixa": summary,
            },
            request,
            competicio.id,
        )
    )


@require_POST
@csrf_protect
def inscripcions_clear_baixa(request, pk):
    competicio = get_object_or_404(Competicio, pk=pk)
    try:
        payload = json.loads(request.body.decode("utf-8"))
    except Exception:
        return HttpResponseBadRequest("JSON invalid")

    try:
        inscripcio_id = int(payload.get("inscripcio_id"))
    except Exception:
        return HttpResponseBadRequest("inscripcio_id invalid")
    inscripcio = get_object_or_404(Inscripcio, pk=inscripcio_id, competicio=competicio)

    before_snapshot = capture_inscripcions_history_snapshot(request, competicio)
    clear_inscripcio_baixa(
        competicio,
        inscripcio,
        user=request.user if getattr(request.user, "is_authenticated", False) else None,
    )
    record_inscripcions_history_entry(
        request,
        competicio,
        action_type="clear_baixa",
        action_label="Treure baixa",
        before_snapshot=before_snapshot,
        after_snapshot=capture_inscripcions_history_snapshot(request, competicio),
    )
    return JsonResponse(
        with_inscripcions_history_payload(
            {"ok": True, "inscripcio_id": inscripcio.id, "baixa": {}},
            request,
            competicio.id,
        )
    )


def inscripcions_baixes_export(request, pk):
    competicio = get_object_or_404(Competicio, pk=pk)
    rows = list(
        active_baixes_qs(competicio)
        .select_related("inscripcio", "comp_aparell", "marcada_per")
        .order_by("inscripcio__nom_i_cognoms", "comp_aparell__ordre", "comp_aparell_id", "id")
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Baixes"
    headers = [
        "Nom i cognoms",
        "Document",
        "Entitat",
        "Categoria",
        "Subcategoria",
        "Abast",
        "Aparell",
        "Motiu",
        "Notes",
        "Data baixa",
        "Marcada per",
    ]
    header_fill = PatternFill("solid", fgColor="FCE4E4")
    for col_idx, label in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for row_idx, baixa in enumerate(rows, start=2):
        ins = baixa.inscripcio
        app = baixa.comp_aparell
        user = baixa.marcada_per
        ws.cell(row=row_idx, column=1, value=getattr(ins, "nom_i_cognoms", "") or "")
        ws.cell(row=row_idx, column=2, value=getattr(ins, "document", "") or "")
        ws.cell(row=row_idx, column=3, value=getattr(ins, "entitat", "") or "")
        ws.cell(row=row_idx, column=4, value=getattr(ins, "categoria", "") or "")
        ws.cell(row=row_idx, column=5, value=getattr(ins, "subcategoria", "") or "")
        ws.cell(row=row_idx, column=6, value="Tota la competicio" if app is None else "Aparell")
        ws.cell(row=row_idx, column=7, value=str(getattr(app, "display_nom", "") or "") if app is not None else "")
        ws.cell(row=row_idx, column=8, value=baixa.motiu or "")
        ws.cell(row=row_idx, column=9, value=baixa.notes or "")
        ws.cell(row=row_idx, column=10, value=baixa.created_at.strftime("%d/%m/%Y %H:%M") if baixa.created_at else "")
        ws.cell(row=row_idx, column=11, value=getattr(user, "username", "") if user is not None else "")

    for col_idx, label in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(14, min(42, len(label) + 8))

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    response = HttpResponse(
        out.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="baixes_competicio_{competicio.pk}.xlsx"'
    return response


__all__ = [
    "InscripcionsListNewView",
    "get_available_table_columns",
    "get_selected_table_columns",
    "inscripcions_baixes_export",
    "inscripcions_clear_baixa",
    "inscripcions_save_birth_year_range_config",
    "inscripcions_save_table_columns",
    "inscripcions_set_baixa",
    "inscripcions_set_aparells",
    "inscripcions_set_group_name",
]
